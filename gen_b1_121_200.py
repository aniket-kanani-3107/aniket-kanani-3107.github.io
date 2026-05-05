# -*- coding: utf-8 -*-
"""
Generate German_Learning_Book_B1_Gujarati_Days_121_200.docx
B1 Course  —  Days 121 to 200
Gujarati + English explanations

This script loads DAYS data from gen_a2_121_200.py (which already
contains full B1-level content) and builds a properly-branded B1 Word book.
"""

import os
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─────────────────────────── colour palette ──────────────────────────────────
BLUE   = RGBColor(0x1A, 0x73, 0xE8)
YELLOW = RGBColor(0xF9, 0xAB, 0x00)
GRAY   = RGBColor(0x55, 0x55, 0x55)
GREEN  = RGBColor(0x0F, 0x96, 0x60)
RED    = RGBColor(0xD9, 0x34, 0x25)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x20, 0x20, 0x20)
PURPLE = RGBColor(0x6A, 0x0D, 0xAD)
ORANGE = RGBColor(0xE6, 0x5C, 0x00)

# ─────────────────────────── helpers ─────────────────────────────────────────

def add_para(doc, text, size_pt=11, bold=False, color=None,
             align=None, space_before=0, space_after=4):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(doc, text, size_pt=11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(size_pt)
    return p


def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def cell_para(cell, text, bold=False, size_pt=10, color=None, align=None):
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.clear()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    return p


def add_vocab_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    headers = ["🇩🇪 German", "🇬🇧 English", "ગુજરાતી", "🔊 Sound like…"]
    for i, h in enumerate(headers):
        set_cell_bg(hdr[i], "1A73E8")
        cell_para(hdr[i], h, bold=True, size_pt=10,
                  color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row().cells
        colors = [DARK, DARK, DARK, GRAY]
        for i, (val, clr) in enumerate(zip(row_data, colors)):
            cell_para(row[i], val, size_pt=10, color=clr)
    doc.add_paragraph()


def add_sentence_table(doc, de, en, gu):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    cells = table.rows[0].cells
    set_cell_bg(cells[0], "E8F0FE")
    set_cell_bg(cells[1], "E6F4EA")
    set_cell_bg(cells[2], "FFF8E1")
    cell_para(cells[0], de, size_pt=10, color=BLUE)
    cell_para(cells[1], en, size_pt=10, color=GREEN)
    cell_para(cells[2], gu, size_pt=10, color=DARK)
    doc.add_paragraph()


def add_tip_box(doc, emoji, label, text, bg="FFF3CD"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, bg)
    cell.paragraphs[0].clear()
    run1 = cell.paragraphs[0].add_run(f"{emoji} {label}  ")
    run1.bold = True
    run1.font.size = Pt(10)
    run1.font.color.rgb = DARK
    run2 = cell.paragraphs[0].add_run(text)
    run2.font.size = Pt(10)
    run2.font.color.rgb = DARK
    doc.add_paragraph()


def add_day(doc, day_num, title, topic, eng_exp, guj_exp,
            vocab, sentences, tips, tasks, mini_test, outcome):
    doc.add_page_break()
    add_para(doc, f"📘  DAY {day_num}", size_pt=21, bold=True, color=BLUE,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2)
    add_para(doc, title, size_pt=14, bold=True, color=YELLOW,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, f"📚 Topic: {topic}", size_pt=10.5, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    doc.add_paragraph()

    add_para(doc, "📖  English Explanation", size_pt=10.5, bold=True,
             color=BLUE, space_after=2)
    add_para(doc, eng_exp, size_pt=10.5, space_after=6)

    add_para(doc, "🇮🇳  ગુજરાતી સ્પષ્ટીકરણ", size_pt=10.5, bold=True,
             color=BLUE, space_after=2)
    add_para(doc, guj_exp, size_pt=10.5, space_after=6)

    add_para(doc, "📋  Vocabulary List", size_pt=10.5, bold=True,
             color=BLUE, space_after=2)
    add_vocab_table(doc, vocab)

    add_para(doc, "💬  Sentence Examples — Real Life", size_pt=10.5,
             bold=True, color=BLUE, space_after=2)
    for de, en, gu in sentences:
        add_sentence_table(doc, de, en, gu)

    for emoji, label, text, bg in tips:
        add_tip_box(doc, emoji, label, text, bg)

    add_para(doc, "✍️  Practice Tasks", size_pt=10.5, bold=True,
             color=BLUE, space_after=2)
    for task in tasks:
        add_bullet(doc, task)
    doc.add_paragraph()

    add_para(doc, "🧪  Mini Test", size_pt=10.5, bold=True,
             color=BLUE, space_after=2)
    for i, (q, a) in enumerate(mini_test, 1):
        add_para(doc, f"Q{i}. {q}", size_pt=10.5, space_after=1)
        add_para(doc, f"   ✅ {a}", size_pt=10.5, color=GREEN, space_after=3)
    doc.add_paragraph()

    add_para(doc, "🏆  Final Outcome", size_pt=10.5, bold=True,
             color=BLUE, space_after=2)
    add_para(doc, outcome, size_pt=10.5, space_after=8)


# ─────────────────────────── load DAYS from A2 script ────────────────────────

def _load_days():
    """
    Read gen_a2_121_200.py, neutralise its build_book() call, and
    return the DAYS list it builds.
    """
    script_dir = os.path.dirname(os.path.abspath(__file__))
    src_path   = os.path.join(script_dir, "gen_a2_121_200.py")
    with open(src_path, encoding="utf-8") as fh:
        src = fh.read()

    # Prevent the source script from running its own build_book()
    # Only replace the bare call at the end (not the def line)
    src = src.replace("\nbuild_book()\n", "\n# build_book() skipped — loaded externally\n")

    ns = {"__name__": "__loaded__"}
    exec(compile(src, src_path, "exec"), ns)   # noqa: S102
    return ns["DAYS"]


# ═══════════════════════════ BUILD B1 BOOK ═══════════════════════════════════

def build_b1_book():
    DAYS = _load_days()

    doc = Document()

    # A4 page setup
    for section in doc.sections:
        section.page_width   = Cm(21)
        section.page_height  = Cm(29.7)
        section.top_margin   = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin  = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ── TITLE PAGE ────────────────────────────────────────────────────────────
    add_para(doc, "🇩🇪", size_pt=72,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=30, space_after=6)

    add_para(doc, "DEUTSCH LERNEN", size_pt=36, bold=True, color=BLUE,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    add_para(doc, "B1 COURSE", size_pt=22, bold=True, color=PURPLE,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    # decorative line
    add_para(doc, "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━",
             size_pt=11, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    add_para(doc, "Days 121 – 200", size_pt=18, bold=True, color=ORANGE,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    add_para(doc, "ગુજરાતી ભાષકો માટે B1 જર્મન અભ્યાસ",
             size_pt=14, bold=True, color=DARK,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    add_para(doc, "For Gujarati Speakers  |  English + Gujarati Explanations",
             size_pt=12, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    add_para(doc, "80 Days  ·  B1 Grammar  ·  Real-Life German  ·  Exam Practice",
             size_pt=11, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    add_para(doc, "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━",
             size_pt=11, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    add_para(doc, "📘  B1 Level  |  Goethe-Institut / telc / ÖSD Preparation",
             size_pt=11, bold=True, color=BLUE,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    add_para(doc, "✍️  Author: Aniket Kanani  |  🗓️  2025",
             size_pt=10, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)

    doc.add_paragraph()

    # ── HOW TO USE ────────────────────────────────────────────────────────────
    doc.add_page_break()
    add_para(doc, "📖  How to Use This Book", size_pt=16, bold=True,
             color=BLUE, space_before=10, space_after=6)
    usage_tips = [
        "📅  Study ONE day at a time — consistency beats speed.",
        "✍️  Write every vocabulary word in a notebook and review it daily.",
        "🗣️  Read every German sentence aloud — B1 demands spoken confidence.",
        "🔁  Revisit previous days each week (spaced repetition works!).",
        "🎯  Complete all practice tasks and mini-tests honestly.",
        "💬  Try to think in German during everyday moments.",
        "🧪  Use the mock tests (Days 187–199) as full B1 exam practice.",
        "🏅  Celebrate your progress — every day you complete is a win!",
    ]
    for t in usage_tips:
        add_bullet(doc, t)
    doc.add_paragraph()

    # ── B1 OVERVIEW ───────────────────────────────────────────────────────────
    add_para(doc, "🟧  B1 COURSE — WHAT YOU WILL MASTER",
             size_pt=14, bold=True, color=GREEN, space_before=8, space_after=6)

    modules = [
        ("Days 121–140", "Core B1 Grammar",
         "Konjunktiv II, passive voice, relative/infinitive clauses, adjective endings, "
         "prepositional verbs, indirect questions, word order (TEKAMOLO)."),
        ("Days 141–150", "Professional German",
         "Job applications, CV language, interviews, workplace communication, "
         "meetings, customer service, salary, contracts, and IT vocabulary."),
        ("Days 151–160", "Society & Daily Life",
         "News, civic life, environment, social media, university, housing, "
         "city administration, travel, and personal finance."),
        ("Days 161–170", "Health & Social Skills",
         "Healthcare system, symptoms, pharmacy, fitness, mental health, "
         "emergencies, family relationships, childcare, and social events."),
        ("Days 171–180", "Communication & Culture",
         "Presentations, debate, negotiation, idioms, formal complaints, "
         "instructions, storytelling, Q&A, and German cultural etiquette."),
        ("Days 181–190", "Exam Skills I",
         "Reading, listening, writing, and speaking strategies; "
         "grammar & vocabulary review; mock tests for each skill."),
        ("Days 191–200", "Exam Skills II & Graduation",
         "Advanced grammar and vocabulary reviews, mock speaking and writing tests, "
         "self-assessment, final mock exam day, and B1 graduation."),
    ]
    for rng, title_, desc in modules:
        add_para(doc, f"📌 {rng}  —  {title_}", size_pt=11, bold=True,
                 color=PURPLE, space_before=4, space_after=1)
        add_para(doc, desc, size_pt=10.5, color=DARK, space_after=4)

    add_para(doc,
             "🎯 After these 80 days your German will be confident, natural, "
             "and ready for the B1 certificate exam.",
             size_pt=11, bold=True, color=BLUE, space_before=6, space_after=4)
    doc.add_paragraph()

    # ── DAILY LESSONS ─────────────────────────────────────────────────────────
    for entry in DAYS:
        add_day(doc, *entry)

    # ── SAVE ──────────────────────────────────────────────────────────────────
    out = "German_Learning_Book_B1_Gujarati_Days_121_200.docx"
    doc.save(out)
    print(f"✅  Saved: {out}")
    print(f"   Total days: {len(DAYS)}")
    return out


if __name__ == "__main__":
    build_b1_book()

#!/usr/bin/env python3
"""Generate German Learning A2 Book (Days 61-120) in Gujarati/English"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

OUTPUT = "/home/runner/work/aniket-kanani-3107.github.io/aniket-kanani-3107.github.io/German_Learning_Book_A2_Gujarati.docx"

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'),'single')
        border.set(qn('w:sz'),'6')
        border.set(qn('w:space'),'0')
        border.set(qn('w:color'),'4472C4')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def heading(doc, text, size=16, bold=True, color='1F3864', align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    r,g,b = tuple(int(color[i:i+2],16) for i in (0,2,4))
    run.font.color.rgb = RGBColor(r,g,b)
    return p

def body(doc, text, size=11, bold=False, color='000000', align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=3):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    r,g,b = tuple(int(color[i:i+2],16) for i in (0,2,4))
    run.font.color.rgb = RGBColor(r,g,b)
    return p

def bullet(doc, text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(size)
    return p

def box_table(doc, text, bg='EBF3FB', text_color='1F3864', bold=True, font_size=11):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = tbl.cell(0,0)
    set_cell_bg(cell, bg)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    r,g,b = tuple(int(text_color[i:i+2],16) for i in (0,2,4))
    run.font.color.rgb = RGBColor(r,g,b)
    doc.add_paragraph()
    return tbl

def sentence_table(doc, german, english, gujarati):
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [Inches(2.2), Inches(2.2), Inches(2.0)]
    colors = ['1F3864','2E5B1E','5C0000']
    bgs = ['D6E4F7','D6F0D6','F7D6D6']
    texts = [german, english, gujarati]
    for i, cell in enumerate(tbl.row_cells(0)):
        set_cell_bg(cell, bgs[i])
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(texts[i])
        run.font.size = Pt(10)
        run.bold = (i==0)
        r,g,b = tuple(int(colors[i][j:j+2],16) for j in (0,2,4))
        run.font.color.rgb = RGBColor(r,g,b)
    doc.add_paragraph()
    return tbl

def vocab_table(doc, rows_data):
    headers = ['🇩🇪 German','🇬🇧 English','ગુજરાતી','🔊 Sound like…']
    tbl = doc.add_table(rows=1+len(rows_data), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.style = 'Table Grid'
    hrow = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        set_cell_bg(cell, '1F3864')
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255,255,255)
    for ri, row_data in enumerate(rows_data):
        row = tbl.rows[ri+1]
        bg = 'EBF3FB' if ri%2==0 else 'FFFFFF'
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(10)
            if ci==0:
                run.bold = True
    doc.add_paragraph()
    return tbl


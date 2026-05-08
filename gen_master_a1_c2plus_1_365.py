# -*- coding: utf-8 -*-
"""Generate a full 365-day German Mastery book (A1 to C2+) for Gujarati speakers."""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x1A, 0x73, 0xE8)
YELLOW = RGBColor(0xF9, 0xAB, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x0F, 0x96, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x20, 0x20, 0x20)
PURPLE = RGBColor(0x6A, 0x0D, 0xAD)
ORANGE = RGBColor(0xE6, 0x5C, 0x00)


def add_para(doc, text, size_pt=11, bold=False, color=None, align=None, space_before=0, space_after=4):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(doc, text, size_pt=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(size_pt)
    return p


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def cell_para(cell, text, bold=False, size_pt=10, color=None, align=None):
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.clear()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    return p


def add_vocab_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    headers = ["🇩🇪 German", "🇬🇧 English", "🇮🇳 Gujarati Meaning", "🔊 Pronunciation"]
    for i, h in enumerate(headers):
        set_cell_bg(hdr[i], "1A73E8")
        cell_para(hdr[i], h, bold=True, size_pt=10, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row().cells
        colors = [DARK, DARK, DARK, GRAY]
        for i, (val, clr) in enumerate(zip(row_data, colors)):
            cell_para(row[i], val, size_pt=9.7, color=clr)
    doc.add_paragraph()


def add_sentence_table(doc, de, en, gu):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    cells = table.rows[0].cells
    set_cell_bg(cells[0], "E8F0FE")
    set_cell_bg(cells[1], "E6F4EA")
    set_cell_bg(cells[2], "FFF8E1")
    cell_para(cells[0], de, size_pt=9.8, color=BLUE)
    cell_para(cells[1], en, size_pt=9.8, color=GREEN)
    cell_para(cells[2], gu, size_pt=9.8, color=DARK)


def add_tip_box(doc, emoji, label, text, bg="FFF3CD"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, bg)
    cell.paragraphs[0].clear()
    run1 = cell.paragraphs[0].add_run(f"{emoji} {label}  ")
    run1.bold = True
    run1.font.size = Pt(10)
    run1.font.color.rgb = DARK
    run2 = cell.paragraphs[0].add_run(text)
    run2.font.size = Pt(10)
    run2.font.color.rgb = DARK
    doc.add_paragraph()


def pick(seq, idx):
    return seq[idx % len(seq)]


def rotate_slice(words, start, size=15):
    out = []
    for i in range(size):
        out.append(words[(start + i) % len(words)])
    return out


LEVELS = [
    {
        "name": "A1",
        "range": (1, 60),
        "subtitle": "Survival German & Confidence",
        "focus": "basic communication, fear removal, and daily survival situations",
        "story": [
            "You arrive at Frankfurt airport and learn your first polite phrases.",
            "You check into a temporary room and learn how to introduce yourself.",
            "You buy groceries and ask simple price questions.",
            "You travel by bus and train using basic direction language.",
            "You visit city offices and fill simple forms with confidence.",
        ],
        "subtopics": [
            "Greetings", "Introducing yourself", "Numbers and prices", "Days and time", "Family words",
            "Food and shopping", "At the station", "At the doctor", "Apartment basics", "Simple past talk",
            "Asking for help", "Phone basics", "Directions", "Weather and clothes", "Restaurant talk",
        ],
        "vocab": [
            ("Hallo", "hello", "હેલો", "HA-lo"), ("Danke", "thank you", "આભાર", "DAN-keh"),
            ("Bitte", "please/you are welcome", "કૃપા કરીને/સ્વાગત", "BIT-teh"), ("Entschuldigung", "sorry", "માફ કરશો", "ent-SHOOL-dee-goong"),
            ("ich", "I", "હું", "ikh"), ("du", "you", "તું", "doo"), ("Sie", "you (formal)", "તમે (આદરથી)", "zee"),
            ("wohnen", "to live", "રહેવું", "VOH-nen"), ("arbeiten", "to work", "કામ કરવું", "AR-bye-ten"),
            ("lernen", "to learn", "શીખવું", "LERN-en"), ("heute", "today", "આજે", "HOY-teh"),
            ("morgen", "tomorrow", "કાલે", "MOR-gen"), ("jetzt", "now", "હવે", "YETST"),
            ("später", "later", "પછી", "SHPAY-ter"), ("Haus", "house", "ઘર", "hows"),
            ("Zimmer", "room", "ખંડ", "TSIM-mer"), ("Bahnhof", "station", "સ્ટેશન", "BAN-hof"),
            ("Supermarkt", "supermarket", "સુપરમાર્કેટ", "ZOO-per-markt"), ("Wasser", "water", "પાણી", "VAS-ser"),
            ("Brot", "bread", "રોટલી/બ્રેડ", "broht"), ("Milch", "milk", "દૂધ", "milkh"), ("Arzt", "doctor", "ડોક્ટર", "artst"),
            ("Termin", "appointment", "મુલાકાત સમય", "ter-MEEN"), ("Straße", "street", "રસ્તો", "SHTRA-seh"),
            ("links", "left", "ડાબે", "links"), ("rechts", "right", "જમણે", "rekhts"), ("geradeaus", "straight", "સીધું", "geh-RAH-deh-ows"),
            ("bezahlen", "to pay", "ચુકવણી કરવી", "beh-TSAH-len"), ("teuer", "expensive", "મોંઘું", "TOY-er"),
            ("günstig", "cheap", "સસ્તુ", "GÜN-stikh"), ("Freund", "friend", "મિત્ર", "froynt"), ("lernen", "learn", "શીખવું", "LERN-en"),
        ],
    },
    {
        "name": "A2",
        "range": (61, 120),
        "subtitle": "Daily Independence",
        "focus": "longer daily conversations, past/future usage, and independent routines",
        "story": [
            "You move into your first apartment and speak with neighbors.",
            "You open a bank account and manage regular appointments.",
            "You handle shopping complaints politely.",
            "You start a part-time role and discuss schedules.",
            "You begin social activities and weekend plans.",
        ],
        "subtopics": [
            "Daily routine", "Past experiences", "Future plans", "Work schedule", "Health habits",
            "Rent and bills", "Bank and money", "Public services", "Travel planning", "Social invitations",
            "Comparisons", "Giving reasons", "Simple opinions", "Problem-solving", "Study habits",
        ],
        "vocab": [
            ("Gewohnheit", "habit", "આદત", "geh-VOHN-hite"), ("früher", "earlier", "પહેલાં", "FRÜ-er"),
            ("künftig", "future", "ભવિષ્યમાં", "KÜNF-tikh"), ("vereinbaren", "arrange", "ગોઠવવું", "fer-EYN-ba-ren"),
            ("Rechnung", "bill", "બીલ", "REKH-noong"), ("Miete", "rent", "ભાડું", "MEE-teh"),
            ("Überweisung", "transfer", "ટ્રાન્સફર", "ü-ber-VY-zung"), ("Konto", "account", "ખાતું", "KON-toh"),
            ("Erfahrung", "experience", "અનુભવ", "er-FAH-roong"), ("vorbereiten", "prepare", "તૈયાર કરવું", "FOR-be-ry-ten"),
            ("besprechen", "discuss", "ચર્ચા કરવી", "beh-SHPREKH-en"), ("verschieben", "postpone", "મુલતવી રાખવું", "fer-SHEE-ben"),
            ("beantragen", "apply for", "અરજી કરવી", "beh-AN-trah-gen"), ("Unterlagen", "documents", "દસ્તાવેજો", "OON-ter-lah-gen"),
            ("regelmäßig", "regularly", "નિયમિત રીતે", "RAY-gel-may-sikh"), ("sparen", "save money", "બચત કરવી", "SHPAH-ren"),
            ("ausgeben", "spend", "ખર્ચ કરવો", "OWS-gay-ben"), ("einladen", "invite", "આમંત્રિત કરવું", "EYN-lah-den"),
            ("absagen", "cancel", "રદ કરવું", "AP-zah-gen"), ("zustimmen", "agree", "સહમત થવું", "TSOO-shtim-men"),
            ("ablehnen", "decline", "નકારવું", "AP-lay-nen"), ("verbessern", "improve", "સુધારવું", "fer-BES-sern"),
            ("fortschritt", "progress", "પ્રગતિ", "FORT-shrit"), ("ziel", "goal", "લક્ષ્ય", "tseel"),
            ("verlässlich", "reliable", "વિશ્વસનીય", "fer-LESS-likh"), ("pünktlich", "punctual", "સમયપાબંદ", "PÜNKT-likh"),
            ("voraussichtlich", "probably", "શક્યતા મુજબ", "for-OWS-zikht-likh"), ("vorschlagen", "suggest", "સૂચવવું", "FOR-shlah-gen"),
            ("erklären", "explain", "સમજાવવું", "er-KLEH-ren"), ("wiederholen", "repeat", "ફરી કહેવું", "VEE-der-ho-len"),
        ],
    },
    {
        "name": "B1",
        "range": (121, 200),
        "subtitle": "Opinion & Storytelling",
        "focus": "structured opinion, narrative clarity, and confident workplace communication",
        "story": [
            "You join team meetings and share ideas.",
            "You handle bureaucratic documents with less stress.",
            "You describe personal stories with details.",
            "You manage conflicts politely with neighbors and colleagues.",
            "You prepare first interview answers in German.",
        ],
        "subtopics": [
            "Giving opinions", "Story sequence", "Workplace communication", "Feedback language", "Conflict phrases",
            "Interview basics", "Education and career", "Media discussion", "Travel incidents", "Service complaints",
            "Goal planning", "Decision making", "Cause and effect", "Emotion language", "Presentation intro",
            "Debate basics", "Polite disagreement", "Personal growth", "Motivation talk", "Community life",
        ],
        "vocab": [
            ("Meiner Meinung nach", "in my opinion", "મારા મત મુજબ", "MY-ner MY-noong nakh"),
            ("einerseits", "on one hand", "એક તરફ", "EYN-er-zites"), ("andererseits", "on the other hand", "બીજી તરફ", "AN-der-er-zites"),
            ("deshalb", "therefore", "એટલા માટે", "DES-halb"), ("dennoch", "nevertheless", "છતાં પણ", "DEN-nokh"),
            ("erzählen", "tell/narrate", "વર્ણન કરવું", "er-TSAY-len"), ("Schilderung", "description", "વર્ણન", "SHIL-de-roong"),
            ("Eindruck", "impression", "છાપ", "EYN-drook"), ("Rückmeldung", "feedback", "પ્રતિસાદ", "RÜK-mel-doong"),
            ("Vereinbarung", "agreement", "સમજૂતી", "fer-EYN-ba-roong"), ("Missverständnis", "misunderstanding", "ગેરસમજ", "MIS-fer-shtent-nis"),
            ("Bewerbung", "job application", "નોખરી અરજી", "beh-VER-boong"), ("Lebenslauf", "CV", "જીવનવૃત્ત", "LAY-bens-lowf"),
            ("Vorstellungsgespräch", "interview", "ઇન્ટરવ્યૂ", "FOR-shtel-loongs-ge-shprekh"), ("Stelle", "position", "પદ", "SHTEL-leh"),
            ("Erfahrung", "experience", "અનુભવ", "er-FAH-roong"), ("Verantwortung", "responsibility", "જવાબદારી", "fer-ANT-vor-toong"),
            ("Leistung", "performance", "પ્રદર્શન", "LY-stoong"), ("verbessern", "improve", "સુધારવું", "fer-BES-sern"),
            ("überzeugen", "convince", "મનાવવું", "ü-ber-TSOY-gen"), ("begründen", "justify", "કારણ આપવું", "beh-GRÜN-den"),
            ("zusammenfassen", "summarize", "સારાંશ આપવો", "tsoo-ZAM-men-fas-sen"), ("Herausforderung", "challenge", "પડકાર", "he-ROWS-for-de-roong"),
            ("Lösung", "solution", "ઉકેલ", "LÖ-zung"), ("Fortbildung", "further training", "વધુ તાલીમ", "FORT-bil-doong"),
            ("zuverlässig", "reliable", "વિશ્વાસપાત્ર", "tsoo-fer-LESS-ikh"), ("zielorientiert", "goal-oriented", "લક્ષ્યમુખી", "tseel-o-ri-en-TEERT"),
            ("offen", "open", "ખુલ્લું", "OF-fen"), ("strukturiert", "structured", "રચિત", "shtruk-too-REERT"),
            ("sachlich", "objective", "તટસ્થ", "ZAKH-likh"),
        ],
    },
    {
        "name": "B2",
        "range": (201, 270),
        "subtitle": "Professional Fluency",
        "focus": "professional speech, nuanced communication, and analytical writing",
        "story": [
            "You join project meetings and write formal updates.",
            "You discuss deadlines, risks, and team coordination.",
            "You present data and answer follow-up questions.",
            "You handle customer complaints and negotiate outcomes.",
            "You build confidence in high-stakes conversations.",
        ],
        "subtopics": [
            "Meeting language", "Project updates", "Formal email style", "Negotiation basics", "Deadline pressure",
            "Presentation flow", "Data interpretation", "Customer communication", "Decision language", "Risk discussion",
            "Team alignment", "Conflict de-escalation", "Process improvement", "Policy explanation", "Professional etiquette",
            "Structured argument", "Complex connectors", "Abstract writing", "Interview excellence", "Career branding",
        ],
        "vocab": [
            ("Abstimmung", "coordination", "સુમેળ", "AP-shtim-moong"), ("Rückfrage", "follow-up question", "પુછપરછ", "RÜK-frah-geh"),
            ("Frist", "deadline", "સમયમર્યાદા", "frist"), ("Umsetzung", "implementation", "અમલીકરણ", "OOM-zet-soong"),
            ("Verhandlung", "negotiation", "વાટાઘાટ", "fer-HAND-loong"), ("Vorgehensweise", "approach", "પદ્ધતિ", "FOR-gay-hens-vy-zeh"),
            ("Ressource", "resource", "સંસાધન", "re-ZOR-seh"), ("Engpass", "bottleneck", "અવરોધ", "ENG-pas"),
            ("Priorität", "priority", "પ્રાથમિકતા", "pri-o-ri-TET"), ("Auswirkung", "impact", "અસર", "OWS-vir-koong"),
            ("Kennzahl", "metric", "માપદંડ", "KEN-tsahl"), ("Dokumentation", "documentation", "દસ્તાવેજીકરણ", "do-ku-men-ta-TSY-ohn"),
            ("Rücksprache", "consultation", "સલાહ ચર્ચા", "RÜK-shprah-kheh"), ("Vereinheitlichung", "standardization", "એકરૂપતા", "fer-EYN-hite-li-khoong"),
            ("Handlungsspielraum", "room for action", "ક્રિયા અવકાશ", "HAND-loongs-shpeel-rowm"), ("verlässlichkeit", "reliability", "વિશ્વસનીયતા", "fer-LESS-likh-kite"),
            ("Koordination", "coordination", "સંકલન", "ko-or-di-na-TSY-ohn"), ("Risikofaktor", "risk factor", "જોખમ પરિબળ", "REE-zi-ko-fak-tor"),
            ("Nachvollziehbar", "understandable", "સમજાય તેવું", "NAKH-fol-tsee-bar"), ("nachhaltig", "sustainable", "ટકાઉ", "NAKH-hal-tikh"),
            ("Rückmeldung", "feedback", "પ્રતિસાદ", "RÜK-mel-doong"), ("Leitfaden", "guideline", "માર્ગદર્શિકા", "LYT-fah-den"),
            ("Anforderung", "requirement", "આવશ્યકતા", "AN-for-de-roong"), ("Freigabe", "approval", "મંજૂરી", "FRY-gah-beh"),
            ("Abweichung", "deviation", "વિચલન", "AP-vy-khoong"), ("Verbesserung", "improvement", "સુધારો", "fer-BES-se-roong"),
            ("Bestandsaufnahme", "status review", "સ્થિતિ સમીક્ષા", "beh-SHTANTS-owf-nah-meh"), ("Ergebnis", "result", "પરિણામ", "er-GAYB-nis"),
            ("Zusammenarbeit", "collaboration", "સહકાર", "tsoo-ZAM-men-ar-bite"), ("Verbindlichkeit", "commitment", "બાધ્યતા", "fer-BIND-likh-kite"),
        ],
    },
    {
        "name": "C1",
        "range": (271, 330),
        "subtitle": "Rhetoric & Advanced Discussion",
        "focus": "argument depth, stylistic control, and persuasive advanced expression",
        "story": [
            "You lead formal discussions with confidence.",
            "You write high-level essays with structure and nuance.",
            "You present difficult topics and manage critical questions.",
            "You adapt your tone for public, academic, and workplace settings.",
            "You mentor junior learners and explain complex ideas clearly.",
        ],
        "subtopics": [
            "Rhetorical precision", "Argument layering", "Subtext awareness", "Counter-argument design", "Context framing",
            "Policy language", "Ethics discussion", "Research interpretation", "Leadership communication", "Public speaking",
            "Register adaptation", "Advanced cohesion", "Critical reading", "Editorial tone", "Synthesis writing",
            "Debate strategy", "Complex rebuttal", "Influence language", "Exam simulation", "Fluency refinement",
        ],
        "vocab": [
            ("Prämisse", "premise", "મૂળ ધારણા", "PRAY-mis-seh"), ("Folgerung", "conclusion", "નિષ્કર્ષ", "FOL-ge-roong"),
            ("Nuancierung", "nuancing", "સૂક્ષ્મતા", "nu-an-SEE-roong"), ("Kohärenz", "coherence", "સુસંગતતા", "ko-hay-RENTS"),
            ("Plausibilität", "plausibility", "વિશ્વસનીયતા", "plow-zi-bi-LEE-tet"), ("Gegenposition", "counter-position", "વિરુદ્ધ સ્થિતિ", "GAY-gen-po-zi-TSY-ohn"),
            ("Einordnung", "contextualization", "સંદર્ભીકરણ", "EYN-ord-noong"), ("Gewichtung", "weighting", "મહત્ત્વ માપ", "geh-VIKH-toong"),
            ("Differenzierung", "differentiation", "વિવેચન", "dif-fe-ren-TSY-roong"), ("Zuspitzung", "sharpening", "તીક્ષ્ણતા", "TSOO-shpit-soong"),
            ("Aussagekraft", "expressive power", "અભિવ્યક્તિ શક્તિ", "OWS-zah-ge-kraft"), ("Verknüpfung", "linkage", "કડી", "fer-KNÜPF-oong"),
            ("Wirkung", "effect", "અસર", "VIR-koong"), ("Schlüssigkeit", "consistency", "તર્કપૂર્ણતા", "SHLÜS-sikh-kite"),
            ("Rahmung", "framing", "ફ્રેમિંગ", "RAH-moong"), ("Einwand", "objection", "આપત્તિ", "EYN-vant"),
            ("Entkräftung", "refutation", "ખંડન", "ent-KREF-toong"), ("Standpunkt", "standpoint", "દૃષ્ટિકોણ", "SHTANT-punkt"),
            ("Abwägung", "consideration", "મૂલ્યાંકન તોલમોલ", "AP-vay-goong"), ("Narrativ", "narrative", "વૃત્તાંત", "na-ra-TEEF"),
            ("Legitimation", "legitimation", "માન્યતા", "le-gi-ti-ma-TSY-ohn"), ("Verhältnismäßigkeit", "proportionality", "અનુપાત યોગ્યતા", "fer-HELTS-nis-may-sikh-kite"),
            ("Zielkonflikt", "goal conflict", "લક્ષ્ય ટકરાવ", "TSEEL-kon-flikt"), ("Nachvollzug", "traceability", "અનુસરણ સ્પષ્ટતા", "NAKH-fol-tsoog"),
            ("Leserführung", "reader guidance", "પાઠક માર્ગદર્શન", "LAY-zer-fü-roong"), ("Sprachökonomie", "language economy", "ભાષા કાર્યક્ષમતા", "SHPRAKH-ö-ko-no-mee"),
            ("Prägnanz", "conciseness", "સંક્ષિપ્ત પ્રભાવ", "PRAYG-nants"), ("Stilbruch", "style break", "શૈલી તૂટણ", "SHTIL-brookh"),
            ("Anschlussfähigkeit", "connectability", "જોડાણક્ષમતા", "AN-shloos-fee-ikh-kite"), ("Reflexion", "reflection", "ચિંતન", "re-FLEK-tsy-ohn"),
        ],
    },
    {
        "name": "C2+",
        "range": (331, 365),
        "subtitle": "Native-like Nuance & Cultural Intelligence",
        "focus": "native-like thought, nuanced register switching, diplomacy, and expressive elegance",
        "story": [
            "You navigate high-level professional and academic circles in German.",
            "You use humor, nuance, and diplomacy naturally.",
            "You mentor others while preserving your own style growth.",
            "You adapt across formal institutions and intimate social contexts.",
            "You complete a full integration journey with linguistic confidence.",
        ],
        "subtopics": [
            "Conceptual precision", "Philosophical nuance", "Humor and irony", "Diplomatic phrasing", "Cultural subtext",
            "Literary interpretation", "Academic authority", "Policy argument", "Public persuasion", "Identity expression",
            "Tone orchestration", "Speech elegance", "Compression mastery", "Counterfactual debate", "Strategic empathy",
            "Intercultural mediation", "Native-like rhythm", "Intellectual dialogue", "Complex storytelling", "Final mastery",
        ],
        "vocab": [
            ("Feinsteuerung", "fine control", "સૂક્ષ્મ નિયંત્રણ", "FYN-shtoy-roong"), ("Tonalität", "tonality", "સ્વરભાવ", "to-na-li-TET"),
            ("Nuance", "nuance", "સૂક્ષ્મ અર્થ", "nu-ANS"), ("Implikation", "implication", "ગૂઢ પરિણામ", "im-pli-ka-TSY-ohn"),
            ("Abstrahierung", "abstraction", "અમૂર્તીકરણ", "ap-stra-HEE-roong"), ("Gedankentiefe", "depth of thought", "વિચાર ઊંડાણ", "geh-DANK-en-TEE-feh"),
            ("Treffsicherheit", "precision accuracy", "ચોક્કસ અસર", "TREF-zikher-hite"), ("Formulierungskunst", "art of phrasing", "વાક્યકળા", "for-moo-LEE-roongs-koonst"),
            ("Souveränität", "mastery/confidence", "પ્રભુત્વ આત્મવિશ્વાસ", "zoo-ve-re-nee-TET"), ("Registerwechsel", "register shift", "રજિસ્ટર ફેરફાર", "reh-GIS-ter-vek-sel"),
            ("Spannungsbogen", "arc of tension", "રોચક વળાંક રેખા", "SHPAN-noongs-bo-gen"), ("Bildsprache", "imagery", "ચિત્રાત્મક ભાષા", "BILT-shprah-kheh"),
            ("Anspielung", "allusion", "ઇશારો", "AN-shpee-loong"), ("Sinnschicht", "layer of meaning", "અર્થસ્તર", "ZIN-shikht"),
            ("Deutungsspielraum", "interpretive space", "અર્થઘટન અવકાશ", "DOY-toongs-shpeel-rowm"), ("Abstufung", "gradation", "સ્તરીકરણ", "AP-shtoo-foong"),
            ("Wendigkeit", "agility", "લવચીકતા", "VEN-dig-kite"), ("Resonanz", "resonance", "પ્રતિધ્વનિ અસર", "re-zo-NANTS"),
            ("Doppelbödigkeit", "double-layered meaning", "દ્વિઅર્થિયતા", "DOP-pel-bö-dig-kite"), ("Unterton", "undertone", "અંતર સ્વર", "OON-ter-tohn"),
            ("Schlagfertigkeit", "quick wit", "તાત્કાલિક બુદ્ધિપ્રતિક્રિયા", "SHLAG-fer-tikh-kite"), ("Selbstkorrektur", "self-correction", "સ્વ-સુધારો", "ZELBST-kor-rek-TOOR"),
            ("Stilsicherheit", "stylistic confidence", "શૈલી વિશ્વાસ", "SHTIL-zikher-hite"), ("Eleganz", "elegance", "આકર્ષક શૈલી", "eh-LE-gants"),
            ("Diskurs", "discourse", "વૈચારિક ચર્ચા", "dis-KOORS"), ("Mehrdeutigkeit", "ambiguity", "બહુઅર્થિયતા", "MAYR-doy-tig-kite"),
            ("Transferleistung", "transfer ability", "રૂપાંતર ક્ષમતા", "trans-FER-ly-stoong"), ("Präzisionsarbeit", "precision work", "ચોક્સાઇ કાર્ય", "pray-tsi-ZY-ohns-ar-bite"),
            ("Sprachintelligenz", "language intelligence", "ભાષિક બુદ્ધિ", "SHPRAKH-in-te-li-GENTS"), ("Weiterentwicklung", "further development", "આગળનો વિકાસ", "VY-ter-ent-vik-loong"),
        ],
    },
]

SCENARIOS = [
    "supermarket", "office", "airport", "doctor", "landlord", "visa office", "university",
    "train station", "restaurant", "dating", "friends meetup", "job interview", "apartment viewing",
]

PRON_GUIDE = [
    "ch → soft friction: keep tongue near upper palate; not hard 'k'.",
    "r → German throat 'r' is softer than Gujarati rolling 'ર'.",
    "ö → shape lips as 'o' but voice 'e'.",
    "ü → shape lips as 'u' but voice 'i'.",
    "ä → between 'એ' and 'ઍ'; keep mouth open and flat.",
    "eu → sounds like 'oy' (as in 'boy').",
    "ei → sounds like 'ai' (as in 'time').",
]

MEMORY_TRICKS = [
    "Create a funny mental image for each new word and place it in your room.",
    "Use a 3-step recall: morning quick review, afternoon usage, night self-test.",
    "Pair each word with one emotion and one physical gesture.",
    "Use story chaining: connect 5 new words in one short visual scene.",
    "Record your own audio flashcards and replay during commute.",
]

MISTAKES = {
    "A1": [
        "Avoid translating word-by-word from Gujarati; keep German verb position simple.",
        "Do not skip articles (der/die/das). They change meaning clarity.",
        "Do not pronounce 'ch' like 'k'; keep it softer.",
    ],
    "A2": [
        "When talking about past events, keep auxiliary + participle structure complete.",
        "Avoid direct English sentence order in subordinate clauses.",
        "Do not overuse present tense when future intention is clear.",
    ],
    "B1": [
        "Avoid overlong sentences without connectors; use deshalb, trotzdem, obwohl.",
        "Differentiate opinion vs fact language clearly.",
        "Keep formal and informal 'you' separate in workplace contexts.",
    ],
    "B2": [
        "Do not use aggressive directness in negotiation emails.",
        "Avoid vague claims; support with one concrete example.",
        "Mind preposition-verb combinations in formal communication.",
    ],
    "C1": [
        "Avoid decorative complexity without logic.",
        "Balance persuasion with evidence and structure.",
        "Check cohesion before submission.",
    ],
    "C2+": [
        "Do not confuse complexity with mastery; precision is higher-level.",
        "Use irony carefully across cultures.",
        "Always adapt register to power-distance and context.",
    ],
}

CULTURE_TIPS = [
    "Punctuality matters deeply; arriving 5 minutes early is appreciated.",
    "In Germany, direct communication is normal; direct does not mean rude.",
    "Separate trash correctly (paper, plastic, bio, residual).",
    "Quiet hours in apartments are respected; avoid loud noise at night.",
    "Formal greetings and clear boundaries build trust faster.",
    "Appointments are valued; cancellations should be communicated early.",
]

PRO_TIPS = [
    "Think in German for 5 minutes after each lesson—no translation allowed.",
    "Shadow one native audio clip daily for pronunciation rhythm.",
    "Use one German-only hour every evening (phone + notes + speech).",
    "Record weekly speaking samples and compare improvements.",
    "Write short daily reflections in German to build natural flow.",
]

LEVEL_MAP = [
    ("A1", 1, 60),
    ("A2", 61, 120),
    ("B1", 121, 200),
    ("B2", 201, 270),
    ("C1", 271, 330),
    ("C2+", 331, 365),
]


def level_for_day(day_num):
    for name, start, end in LEVEL_MAP:
        if start <= day_num <= end:
            return name, start, end
    raise ValueError(f"Invalid day: {day_num}")


def make_day(level_cfg, day_num):
    level = level_cfg["name"]
    _, start, _ = level_for_day(day_num)
    local_index = day_num - start

    subtopic = pick(level_cfg["subtopics"], local_index)
    scenario = pick(SCENARIOS, local_index)
    story_line = pick(level_cfg["story"], local_index)
    pron = pick(PRON_GUIDE, local_index)
    memory = pick(MEMORY_TRICKS, local_index)
    mistake = pick(MISTAKES[level], local_index)
    culture = pick(CULTURE_TIPS, local_index)
    pro = pick(PRO_TIPS, local_index)

    vocab = rotate_slice(level_cfg["vocab"], local_index * 2, size=15)
    k1, k2, k3 = vocab[0][0], vocab[1][0], vocab[2][0]

    title = f"{subtopic.upper()} 🚀"
    topic = f"{level} — {subtopic}"

    today_you_will = (
        f"By the end of Day {day_num}, you will confidently use {subtopic.lower()} language in a real-life {scenario} context, "
        f"speak with better pronunciation control, and remember at least 15 useful words actively."
    )

    eng_exp = (
        f"Today is about {subtopic.lower()}. Keep it simple: understand the pattern, say it out loud, and use it in a real situation. "
        f"You are building {level_cfg['focus']}. Story mode: {story_line} "
        f"Do not fear mistakes—every correction today makes your German stronger tomorrow."
    )

    guj_exp = (
        f"આજે આપણે {subtopic.lower()} પર કામ કરીશું. સરળ રીતે શીખો: નિયમ સમજો, ઊંચે અવાજે બોલો અને વાસ્તવિક પરિસ્થિતિમાં વાપરો. "
        f"તમે {level_cfg['focus']} વિકસાવી રહ્યા છો. સ્ટોરી મોડ: {story_line} "
        f"ભૂલથી ડરશો નહીં—આજનો દરેક સુધારો તમારું જર્મન કાલે વધુ મજબૂત બનાવશે."
    )

    dialogues = [
        (
            f"Ich brauche Hilfe im {scenario}.",
            f"I need help at the {scenario}.",
            f"મને {scenario} માં મદદ જોઈએ છે."
        ),
        (
            f"Können Sie mir kurz erklären, wie {k1} hier funktioniert?",
            f"Could you briefly explain how {k1} works here?",
            f"શું તમે ટૂંકમાં સમજાવી શકો કે અહીં {k1} કેવી રીતે કામ કરે છે?"
        ),
        (
            f"Danke, mit {k2} und {k3} kann ich es jetzt besser sagen.",
            f"Thanks, with {k2} and {k3}, I can say it better now.",
            f"આભાર, {k2} અને {k3} સાથે હું હવે વધુ સારી રીતે કહી શકું છું."
        ),
        (
            "Kein Problem, Schritt für Schritt wird es leichter.",
            "No problem, step by step it becomes easier.",
            "કોઈ સમસ્યા નથી, પગલું પગલું આગળ વધતાં સરળ બને છે."
        ),
        (
            "Ich übe das heute dreimal laut.",
            "I will practice this aloud three times today.",
            "હું આજે આ ત્રણ વાર ઊંચે અવાજે પ્રેક્ટિસ કરીશ."
        ),
    ]

    mini_test = [
        ("Translate into German: 'I can handle this situation calmly.'", "Ich kann diese Situation ruhig bewältigen."),
        ("Fill in: Ich _____ heute gezielt Deutsch. (lernen)", "lerne"),
        ("Build one sentence with today's word.", vocab[0][0]),
        ("Listening imagination: What phrase would you use first in this scenario?", "Use a polite opener + your request clearly."),
        ("Speaking prompt: Say a 30-second response about today's topic.", "Clear intro + one example + confident closing."),
    ]

    outcome = (
        f"After today, you can now manage {subtopic.lower()} conversations in {scenario} with better confidence, clearer pronunciation, "
        f"and smarter word recall."
    )

    return {
        "day": day_num,
        "title": title,
        "topic": topic,
        "today": today_you_will,
        "eng": eng_exp,
        "guj": guj_exp,
        "vocab": vocab,
        "dialogues": dialogues,
        "pron": pron,
        "memory": memory,
        "mistake": mistake,
        "culture": culture,
        "pro": pro,
        "tasks": [
            f"Writing: Write 8–12 lines on '{subtopic}' using at least 6 new words.",
            f"Speaking: Record a 2–4 minute response for a {scenario} situation.",
            "Listening: Watch/listen to a short German clip and note 5 phrases you hear.",
            "Shadowing: Repeat one native audio clip 5 times with matching rhythm.",
            "Self-record: Compare first and last recording to notice pronunciation improvement.",
        ],
        "mini_test": mini_test,
        "outcome": outcome,
    }


def add_day(doc, data):
    doc.add_page_break()

    add_para(doc, f"📘 DAY {data['day']}", size_pt=21, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2)
    add_para(doc, data["title"], size_pt=14, bold=True, color=YELLOW, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, f"📚 Topic: {data['topic']}", size_pt=10.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    add_para(doc, "🎯 Today You Will:", size_pt=11, bold=True, color=BLUE, space_after=2)
    add_para(doc, data["today"], size_pt=10.5, space_after=6)

    add_para(doc, "📖 English Explanation", size_pt=11, bold=True, color=BLUE, space_after=2)
    add_para(doc, data["eng"], size_pt=10.4, space_after=6)

    add_para(doc, "🇮🇳 ગુજરાતી સ્પષ્ટીકરણ", size_pt=11, bold=True, color=BLUE, space_after=2)
    add_para(doc, data["guj"], size_pt=10.4, space_after=6)

    add_para(doc, "📋 Vocabulary Section", size_pt=11, bold=True, color=BLUE, space_after=2)
    add_vocab_table(doc, data["vocab"])

    add_para(doc, "💬 Real-Life Conversation Examples", size_pt=11, bold=True, color=BLUE, space_after=2)
    for de, en, gu in data["dialogues"]:
        add_sentence_table(doc, de, en, gu)
    doc.add_paragraph()

    add_tip_box(doc, "🔊", "Pronunciation Training", data["pron"], bg="E8F0FE")
    add_tip_box(doc, "🧠", "Memory System", data["memory"], bg="FFF8E1")
    add_tip_box(doc, "⚠️", "Common Mistake", data["mistake"], bg="FDECEA")
    add_tip_box(doc, "🇩🇪", "German Culture Tip", data["culture"], bg="E6F4EA")
    add_tip_box(doc, "💡", "Pro Tip", data["pro"], bg="F3E8FD")

    add_para(doc, "✍️ Practice Tasks", size_pt=11, bold=True, color=BLUE, space_after=2)
    for task in data["tasks"]:
        add_bullet(doc, task)
    doc.add_paragraph()

    add_para(doc, "🧪 Mini Test", size_pt=11, bold=True, color=BLUE, space_after=2)
    for i, (q, a) in enumerate(data["mini_test"], 1):
        add_para(doc, f"Q{i}. {q}", size_pt=10.2, space_after=1)
        add_para(doc, f"✅ {a}", size_pt=10.2, color=GREEN, space_after=3)

    add_para(doc, "🏆 Final Outcome", size_pt=11, bold=True, color=BLUE, space_after=2)
    add_para(doc, data["outcome"], size_pt=10.5, space_after=8)


def add_level_exam_pack(doc, level_name, start_day, end_day):
    doc.add_page_break()
    add_para(doc, f"🧪 {level_name} LEVEL EXAM SIMULATION PACK", size_pt=16, bold=True, color=PURPLE, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=4)
    add_para(doc, f"Covers Days {start_day}-{end_day} | Goethe/telc/ÖSD/TestDaF style practice", size_pt=10.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    for exam_no in range(1, 11):
        add_para(doc, f"📘 Full Exam {exam_no}", size_pt=12, bold=True, color=BLUE, space_before=4, space_after=2)
        add_bullet(doc, "Reading: 2 texts + 8 comprehension questions")
        add_bullet(doc, "Writing: 1 guided task + 1 free-response task")
        add_bullet(doc, "Listening: 2 scenarios with key-detail extraction")
        add_bullet(doc, "Speaking: 2-minute intro + 4-minute discussion task")
        add_bullet(doc, "Vocabulary & Grammar: 20 applied items from the level range")
        add_bullet(doc, "Real simulation timing: strict pacing + self-evaluation rubric")

        add_tip_box(
            doc,
            "🎯",
            "Exam Goal",
            f"Finish Exam {exam_no} with calm pacing, clear structure, and confident pronunciation. Focus on accuracy + fluency + strategy.",
            bg="E8F0FE",
        )


def build_days():
    items = []
    for cfg in LEVELS:
        start, end = cfg["range"]
        for day in range(start, end + 1):
            items.append((cfg, make_day(cfg, day)))
    return items


def build_book():
    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    add_para(doc, "🇩🇪", size_pt=60, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=20, space_after=6)
    add_para(doc, "GERMAN MASTERY PROGRAM", size_pt=30, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
    add_para(doc, "A1 → C2+ | 365-DAY ROADMAP", size_pt=18, bold=True, color=PURPLE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
    add_para(doc, "ગુજરાતી + English Support", size_pt=14, bold=True, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_para(doc, "Daily Life · Work · Exams · Culture · Fluency", size_pt=11, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    add_para(doc, "📖 How to Use This Master Book", size_pt=14, bold=True, color=BLUE, space_after=4)
    for t in [
        "Study 5–6 focused hours daily with active speaking.",
        "Follow the streak: one day at a time, no skipping practice tasks.",
        "Use self-recording daily for pronunciation and confidence growth.",
        "Revise weekly with memory loops and mini tests.",
        "Complete each level exam pack before moving up.",
    ]:
        add_bullet(doc, t)

    current_level = None
    level_start = None
    for cfg, day_data in build_days():
        level_name = cfg["name"]
        start, end = cfg["range"]

        if level_name != current_level:
            current_level = level_name
            level_start = start
            doc.add_page_break()
            add_para(doc, f"🟦 LEVEL {level_name} — {cfg['subtitle']}", size_pt=16, bold=True, color=PURPLE, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=3)
            add_para(doc, f"Days {start} to {end}", size_pt=12, bold=True, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
            add_para(doc, f"Level Focus: {cfg['focus']}", size_pt=10.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

        add_day(doc, day_data)

        if day_data["day"] == end:
            add_level_exam_pack(doc, level_name, level_start, end)

    out = "German_Learning_Book_A1_to_C2plus_Gujarati_Days_1_365.docx"
    doc.save(out)
    print(f"✅ Saved: {out}")
    print("   Days: 365")


if __name__ == "__main__":
    build_book()

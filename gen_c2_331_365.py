# -*- coding: utf-8 -*-
"""Generate German_Learning_Book_C2_Gujarati_Days_331_365.docx (Days 331-365)"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x1A, 0x73, 0xE8)
YELLOW = RGBColor(0xF9, 0xAB, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x0F, 0x96, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x20, 0x20, 0x20)
PURPLE = RGBColor(0x6A, 0x0D, 0xAD)
ORANGE = RGBColor(0xE6, 0x5C, 0x00)


def add_para(doc, text, size_pt=11, bold=False, color=None, align=None, space_before=0, space_after=4):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(doc, text, size_pt=11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(size_pt)
    return p


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def cell_para(cell, text, bold=False, size_pt=10, color=None, align=None):
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.clear()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    return p


def add_vocab_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    headers = ["🇩🇪 German", "🇬🇧 English", "ગુજરાતી", "🔊 Sound like…"]
    for i, h in enumerate(headers):
        set_cell_bg(hdr[i], "1A73E8")
        cell_para(hdr[i], h, bold=True, size_pt=10, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row().cells
        colors = [DARK, DARK, DARK, GRAY]
        for i, (val, clr) in enumerate(zip(row_data, colors)):
            cell_para(row[i], val, size_pt=10, color=clr)
    doc.add_paragraph()


def add_sentence_table(doc, de, en, gu):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    cells = table.rows[0].cells
    set_cell_bg(cells[0], "E8F0FE")
    set_cell_bg(cells[1], "E6F4EA")
    set_cell_bg(cells[2], "FFF8E1")
    cell_para(cells[0], de, size_pt=10, color=BLUE)
    cell_para(cells[1], en, size_pt=10, color=GREEN)
    cell_para(cells[2], gu, size_pt=10, color=DARK)
    doc.add_paragraph()


def add_tip_box(doc, emoji, label, text, bg="FFF3CD"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, bg)
    cell.paragraphs[0].clear()
    run1 = cell.paragraphs[0].add_run(f"{emoji} {label}  ")
    run1.bold = True
    run1.font.size = Pt(10)
    run1.font.color.rgb = DARK
    run2 = cell.paragraphs[0].add_run(text)
    run2.font.size = Pt(10)
    run2.font.color.rgb = DARK
    doc.add_paragraph()


def add_day(doc, day_num, title, topic, eng_exp, guj_exp, vocab, sentences, tips, tasks, mini_test, outcome):
    doc.add_page_break()
    add_para(doc, f"📘  DAY {day_num}", size_pt=21, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2)
    add_para(doc, title, size_pt=14, bold=True, color=YELLOW, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, f"📚 Topic: {topic}", size_pt=10.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    doc.add_paragraph()

    add_para(doc, "📖  English Explanation", size_pt=10.5, bold=True, color=BLUE, space_after=2)
    add_para(doc, eng_exp, size_pt=10.5, space_after=6)

    add_para(doc, "🇮🇳  ગુજરાતી સ્પષ્ટીકરણ", size_pt=10.5, bold=True, color=BLUE, space_after=2)
    add_para(doc, guj_exp, size_pt=10.5, space_after=6)

    add_para(doc, "📋  Vocabulary List", size_pt=10.5, bold=True, color=BLUE, space_after=2)
    add_vocab_table(doc, vocab)

    add_para(doc, "💬  Sentence Examples — Real Life", size_pt=10.5, bold=True, color=BLUE, space_after=2)
    for de, en, gu in sentences:
        add_sentence_table(doc, de, en, gu)

    for emoji, label, text, bg in tips:
        add_tip_box(doc, emoji, label, text, bg)

    add_para(doc, "✍️  Practice Tasks", size_pt=10.5, bold=True, color=BLUE, space_after=2)
    for task in tasks:
        add_bullet(doc, task)
    doc.add_paragraph()

    add_para(doc, "🧪  Mini Test", size_pt=10.5, bold=True, color=BLUE, space_after=2)
    for i, (q, a) in enumerate(mini_test, 1):
        add_para(doc, f"Q{i}. {q}", size_pt=10.5, space_after=1)
        add_para(doc, f"   ✅ {a}", size_pt=10.5, color=GREEN, space_after=3)
    doc.add_paragraph()

    add_para(doc, "🏆  Final Outcome", size_pt=10.5, bold=True, color=BLUE, space_after=2)
    add_para(doc, outcome, size_pt=10.5, space_after=8)


MODULES = [
    {
        "range": (331, 337),
        "name": "Philosophical Argumentation and Abstract Precision",
        "focus": "conceptual rigor, epistemic nuance, and layered abstract reasoning",
        "vocab": [
            ("die Begriffsarbeit", "conceptual work", "સંકલ્પનાત્મક કાર્ય", "beh-GRIFS-ar-bite"),
            ("der Gedankengang", "line of thought", "વિચાર પ્રવાહ", "geh-DANK-en-gang"),
            ("die Erkenntnisschranke", "limit of knowledge", "જ્ઞાનની મર્યાદા", "er-KENT-nis-shran-ke"),
            ("die Plausibilisierung", "plausibility building", "વિશ્વસનીય બનાવવાની પ્રક્રિયા", "plow-zi-bi-li-ZEE-roong"),
            ("die Vorannahme", "prior assumption", "પૂર્વધારણા", "FOR-an-na-me"),
            ("die Denkfigur", "thought pattern", "વિચાર માળખું", "DENK-fi-goor"),
            ("die Gegenintuition", "counter-intuition", "વિરુદ્ધ આંતરિક અનુમાન", "GAY-gen-in-too-ee-TSY-ohn"),
            ("die Abstraktionsebene", "level of abstraction", "અમૂર્તતા સ્તર", "ap-strak-TSY-ohs-EH-be-ne"),
            ("die Schlusslogik", "inferential logic", "નિષ્કર્ષ તર્ક", "SHLOOS-lo-gik"),
            ("die Tragweite", "scope/implication", "અસર વ્યાપકતા", "TRAHG-vy-teh"),
            ("die Denkvoraussetzung", "cognitive precondition", "વિચાર પૂર્વશરત", "DENK-fo-rows-zet-soong"),
            ("die Unschärfe", "blur/indeterminacy", "અસ્પષ્ટતા", "OON-sher-feh"),
            ("die Reflexionsstufe", "level of reflection", "ચિંતન સ્તર", "re-flex-TSY-ohns-shtoo-feh"),
            ("die Konsistenzprüfung", "consistency check", "સુસંગતતા ચકાસણી", "kon-zis-TENTS-prü-foong"),
            ("die Implikationskette", "chain of implications", "પરિણામ શ્રેણી", "im-pli-ka-TSY-ohns-ket-te"),
            ("das Deutungsmuster", "interpretive pattern", "અર્થઘટન નમૂનો", "DOY-toongs-moo-ster"),
            ("die Selbstrelativierung", "self-relativization", "સ્વ-સાપેક્ષતા", "ZELBST-re-la-ti-VEE-roong"),
            ("die Urteilsenthaltung", "suspension of judgment", "નિર્ણય અટકાવવો", "OOR-tyls-ent-hal-toong"),
            ("der Präzisierungsbedarf", "need for clarification", "સ્પષ્ટીકરણની જરૂર", "pray-tsi-ZEE-roongs-be-darf"),
            ("die Abwägungstiefe", "depth of evaluation", "મૂલ્યાંકનની ઊંડાણ", "AP-vay-goongs-TEE-feh"),
        ],
    },
    {
        "range": (338, 344),
        "name": "Literary Interpretation and Cultural Critique",
        "focus": "symbolic reading, tonal sensitivity, and historically informed interpretation",
        "vocab": [
            ("die Erzählinstanz", "narrative voice", "વર્ણન અવાજ", "er-TSAYL-in-stants"),
            ("die Mehrdeutigkeitsschicht", "layer of ambiguity", "બહુઅર્થિય સ્તર", "MAYR-doy-tig-kites-shikht"),
            ("die Motivführung", "motif development", "મોટિફ પ્રવાહ", "mo-TEEF-fü-roong"),
            ("der Bedeutungshorizont", "horizon of meaning", "અર્થ ક્ષિતિજ", "beh-DOY-toongs-ho-ri-tsont"),
            ("die Lesart", "interpretation/reading", "વાચન અર્થ", "LAY-zart"),
            ("die Brechung", "ironic break/refraction", "વિખંડિત વળાંક", "BRE-khoong"),
            ("die Tiefenstruktur", "deep structure", "આંતરિક રચના", "TEE-fen-shtruk-toor"),
            ("die Verweisstruktur", "network of references", "સંદર્ભ રચના", "fer-VICE-shtruk-toor"),
            ("die Zeitgebundenheit", "historical situatedness", "સમયબંધ સ્થિતિ", "TSYTE-ge-boon-den-hite"),
            ("die Rezeptionsgeschichte", "reception history", "ગ્રાહ્ય ઇતિહાસ", "re-tsep-TSY-ohns-ge-shikh-teh"),
            ("die Stimmungsführung", "mood management", "માહોલ નિર્માણ", "SHTIM-moongs-fü-roong"),
            ("die Perspektivverschiebung", "perspective shift", "દૃષ્ટિકોણ ફેરફાર", "per-spek-TEEF-fer-shee-boong"),
            ("die Codierung", "coding", "સંકેતન", "ko-DEE-roong"),
            ("die Entzifferung", "decoding", "સંકેત ઉકેલવું", "ent-TSIF-fe-roong"),
            ("die Textur", "texture", "લેખીય સપાટી/ટેક્સ્ચર", "teks-TOOR"),
            ("die Intertextualität", "intertextuality", "અંતર-પાઠીયતા", "in-ter-teks-too-a-li-TET"),
            ("die Kanondebatte", "canon debate", "સાહિત્ય કાનન ચર્ચા", "ka-NOHN-de-BAT-teh"),
            ("die Deutungsoffenheit", "interpretive openness", "અર્થઘટનની ખુલ્લાશ", "DOY-toongs-of-fen-hite"),
            ("die Stilspur", "stylistic trace", "શૈલી છાપ", "SHTIL-shpoor"),
            ("die Resonanzfläche", "resonance surface", "પ્રતિધ્વનિ ક્ષેત્ર", "re-zo-NANTS-fle-kheh"),
        ],
    },
    {
        "range": (345, 351),
        "name": "Diplomacy, Mediation, and Institutional Language",
        "focus": "consensus engineering, calibrated politeness, and high-stakes institutional communication",
        "vocab": [
            ("die Formulierungsreserve", "reserve in wording", "શબ્દપ્રયોગ સાવચેતી", "for-moo-LEE-roongs-re-zer-ve"),
            ("die Konsensarchitektur", "consensus architecture", "સહમતિ રચના", "kon-ZENS-ar-khi-tek-TOOR"),
            ("die Annäherungsformel", "formula of rapprochement", "નજીક લાવતી રચના", "an-NAY-he-roongs-for-mel"),
            ("die Vertraulichkeitsstufe", "level of confidentiality", "ગોપનીયતા સ્તર", "fer-TROW-likh-kites-shtoo-feh"),
            ("die Interessenlage", "interest position", "હિત સ્થિતિ", "in-te-RESS-en-la-geh"),
            ("die Rückversicherung", "reassurance", "પુનઃખાતરી", "RÜK-fer-zikhe-roong"),
            ("der Vorbehalt", "reservation", "આપત્તિ/અનામત", "FOR-be-halt"),
            ("die Sprachregelung", "agreed wording", "નક્કી કરેલો શબ્દપ્રયોગ", "SHPRAKH-ray-ge-loong"),
            ("die Kompromisslinie", "compromise line", "સમાધાન રેખા", "kom-PRO-mis-lee-nee"),
            ("die Vermittlungsrolle", "mediating role", "મધ્યસ્થી ભૂમિકા", "fer-MIT-loongs-rol-leh"),
            ("die Gesprächsdisziplin", "discussion discipline", "ચર્ચા શિસ્ત", "ge-SHPREKHS-di-si-PLIN"),
            ("die Beschlussreife", "decision readiness", "નિર્ણય તૈયારપણું", "beh-SHLOOS-ry-feh"),
            ("die Eskalationsvermeidung", "escalation avoidance", "ઉત્તેજન ટાળવું", "es-ka-la-TSY-ohns-fer-MY-doong"),
            ("die Formulierungshoheit", "authority over wording", "શબ્દપ્રયોગ પર નિયંત્રણ", "for-moo-LEE-roongs-ho-heit"),
            ("die Sondierung", "exploratory sounding", "પ્રારંભિક તપાસ", "zon-DEE-roong"),
            ("die Anschlussfähigkeit", "capacity to connect", "સંદર્ભ જોડાણક્ષમતા", "AN-shloos-fee-ikh-kite"),
            ("die Positionsangleichung", "position alignment", "સ્થિતિ સુમેળ", "po-zi-TSY-ohns-an-gly-khoong"),
            ("die Belastungsprobe", "stress test", "દબાણ પરીક્ષા", "be-LAST-oongs-pro-beh"),
            ("die Verhandlungsmarge", "negotiation margin", "વાટાઘાટ અવકાશ", "fer-HAND-loongs-mar-zhe"),
            ("die Schlussformel", "closing formula", "સમાપન સૂત્ર", "SHLOOS-for-mel"),
        ],
    },
    {
        "range": (352, 358),
        "name": "Research Publication and Expert Debate",
        "focus": "disciplinary precision, publication-ready argumentation, and expert-level rebuttal control",
        "vocab": [
            ("die Begutachtung", "peer review", "સમીક્ષા મૂલ્યાંકન", "beh-GOO-takh-toong"),
            ("die Anschlussfrage", "follow-up research question", "આગળનો સંશોધન પ્રશ્ન", "AN-shloos-fra-geh"),
            ("die Fachterminologie", "specialist terminology", "વિષયક પરિભાષા", "FAKH-ter-mi-no-lo-GEE"),
            ("die Einbettung", "embedding/contextual embedding", "સંદર્ભમાં ગોઠવણ", "EYN-bet-toong"),
            ("die Einordnungstiefe", "depth of contextualization", "સંદર્ભ ઊંડાણ", "EYN-ord-noongs-TEE-feh"),
            ("die Befundlage", "state of findings", "નિષ્કર્ષ સ્થિતિ", "be-FOOND-la-geh"),
            ("die Gegenlektüre", "counter-reading", "પ્રતિવાંચન", "GAY-gen-lek-TÜ-reh"),
            ("die Replizierbarkeit", "replicability", "પુનરાવર્તિત ક્ષમતા", "re-pli-tseer-BAR-kite"),
            ("die Methodenkritik", "method critique", "પદ્ધતિ વિવેચન", "meh-TOH-den-kri-tik"),
            ("die Evidenzabstufung", "evidence gradation", "પુરાવા સ્તરીકરણ", "eh-vee-DENTS-ap-shtoo-foong"),
            ("die Fachdebatte", "specialist debate", "વિશેષજ્ઞ ચર્ચા", "FAKH-de-BAT-teh"),
            ("die Präregistrierung", "pre-registration", "પૂર્વ નોંધણી", "pray-re-gis-TRIE-roong"),
            ("die Ergebnisschärfung", "result sharpening", "પરિણામ તીક્ષ્ણતા", "er-GAYB-nis-sher-foong"),
            ("die Einwandbehandlung", "handling objections", "આપત્તિ સંભાળ", "EYN-vant-be-hand-loong"),
            ("die Disziplingrenze", "disciplinary boundary", "વિષય મર્યાદા", "di-tsi-PLIN-grent-seh"),
            ("die Transferleistung", "transfer achievement", "સ્થાનાંતર પ્રયોગ શક્તિ", "trans-FER-ly-stoong"),
            ("die Vorveröffentlichung", "preprint/publication ahead", "પૂર્વ પ્રકાશન", "FOR-fer-of-fent-li-khoong"),
            ("die Ergebnisnarration", "narration of results", "પરિણામ વૃત્તાંત", "er-GAYB-nis-na-ra-TSY-ohn"),
            ("die Fachresonanz", "disciplinary resonance", "વિષયક પ્રતિસાદ", "FAKH-re-zo-NANTS"),
            ("die Publikationsreife", "publication readiness", "પ્રકાશન તૈયારી", "pu-bli-ka-TSY-ohs-ry-feh"),
        ],
    },
    {
        "range": (359, 365),
        "name": "C2 Mastery, Stylistic Command, and Lifelong Refinement",
        "focus": "register orchestration, tonal precision, and integrated near-native performance",
        "vocab": [
            ("die Registerfeinsteuerung", "fine register control", "રજિસ્ટરની સૂક્ષ્મ નિયંત્રણ", "reh-GIS-ter-fyn-shtoy-roong"),
            ("die Tonalitätssteuerung", "tonality control", "સ્વર નિયંત્રણ", "to-na-li-TETS-shtoy-roong"),
            ("die Ausdrucksdisziplin", "discipline of expression", "અભિવ્યક્તિ શિસ્ત", "OWS-drooks-di-si-PLIN"),
            ("die Verdichtungsleistung", "compression performance", "ઘનતા સર્જન શક્તિ", "fer-DIKH-toongs-ly-stoong"),
            ("die Nuancenarbeit", "nuance work", "સૂક્ષ્મતા પર કામ", "nu-AN-sen-ar-bite"),
            ("die Souveränitätsmarke", "marker of mastery", "નિષ્ણાતતા ચિહ્ન", "zoo-ve-re-nee-TETS-mar-keh"),
            ("die Idiombeherrschung", "idiom mastery", "રૂઢિપ્રયોગ કાબૂ", "ee-dee-OHM-be-herr-shoong"),
            ("die Redeökonomie", "economy of speech", "વાણી કાર્યક્ષમતા", "RAY-deh-ö-ko-no-mee"),
            ("die Prägnanzlinie", "line of concise force", "સંક્ષિપ્ત અસર રેખા", "PRAYG-nants-lee-nee"),
            ("die Selbstrevision", "self-revision", "સ્વ-સંપાદન", "ZELBST-re-vi-ZY-ohn"),
            ("die Sprecherhaltung", "speaker stance", "વક્તા વલણ", "SHPRE-kher-hal-toong"),
            ("die Formulierungsagilität", "agility in phrasing", "વાક્યરચના લવચીકતા", "for-moo-LEE-roongs-a-gi-li-TET"),
            ("die Kontextsensibilität", "context sensitivity", "સંદર્ભ સંવેદનશીલતા", "kon-TEKST-zen-si-bi-li-TET"),
            ("die Eleganzreserve", "reserve of elegance", "શૈલીપૂર્ણ સંગ્રહ", "eh-LE-gants-re-zer-ve"),
            ("die Sprechverdichtung", "condensed speaking", "ઘનતાપૂર્ણ બોલચાલ", "SHPREKH-fer-DIKH-toong"),
            ("die Korrektursicherheit", "correction confidence", "સુધારા વિશ્વાસ", "kor-REK-toor-zikher-hite"),
            ("die Stilsicherheit", "stylistic confidence", "શૈલી વિશ્વાસ", "SHTIL-zikher-hite"),
            ("die Transferpräzision", "precision in transfer", "રૂપાંતર ચોકસાઈ", "trans-FER-pray-tsi-ZY-ohn"),
            ("die Abschlusskompetenz", "finishing competence", "સમાપન કુશળતા", "AP-shloos-kom-pe-TENTS"),
            ("die Weiterentwicklung", "continued development", "આગળનો વિકાસ", "VY-ter-ent-vik-loong"),
        ],
    },
]

SUBTOPICS = [
    "Conceptual framing at C2", "Epistemic caution and certainty limits", "Implicit assumptions under scrutiny", "Multi-perspective weighing", "Contradiction management", "Thesis refinement under pressure", "Philosophical synthesis",
    "Narrative voice and distance", "Symbolic layering", "Ambiguity as meaning", "Historical resonance in texts", "Stylistic texture", "Intertextual bridges", "Cultural critique through literature",
    "Diplomatic softening", "Consensus architecture", "Strategic concessions", "Confidential briefing language", "Multilateral positioning", "Crisis de-escalation talk", "Formal resolution drafting",
    "Peer-review stance", "Methodological caveats", "Disciplinary framing", "Rebuttal writing", "Data narration for experts", "Conference intervention", "Publication polish",
    "Register orchestration", "Tonal precision", "Idiomatic restraint", "Compression under pressure", "Oral agility at speed", "Final integrated simulation", "C2 graduation and refinement",
]


def pick_vocab(bank, start, size=15, context=""):
    if size > len(bank):
        detail = context or "unknown context"
        raise ValueError(
            f"Vocabulary bank size mismatch for {detail}: requested {size} words but only {len(bank)} available."
        )
    out = []
    for i in range(size):
        out.append(bank[(start + i) % len(bank)])
    return out


def build_day_entry(day_num, module, subtopic, idx):
    vocab = pick_vocab(module["vocab"], idx, context=f"Day {day_num} ({module['name']})")
    k1, k2, k3 = vocab[0][0], vocab[1][0], vocab[2][0]

    title = f"{subtopic.upper()} 🚀"
    topic = f"{module['name']} — {subtopic}"

    eng_exp = (
        f"Day {day_num} moves your German into C2 territory through '{subtopic}'. "
        f"Today you sharpen {module['focus']}. The aim is not just correctness, but intellectual authority, "
        f"stylistic flexibility, and effortless adaptation across expert, cultural, and public-facing contexts."
    )
    guj_exp = (
        f"Day {day_num} માં તમે '{subtopic}' દ્વારા C2 સ્તર તરફ આગળ વધો છો. આજે તમારું ધ્યાન "
        f"{module['focus']} પર છે. હેતુ માત્ર સાચું જર્મન નહીં, પરંતુ બૌદ્ધિક અસર, શૈલીની લવચીકતા "
        f"અને નિષ્ણાત, સાંસ્કૃતિક તથા જાહેર સંદર્ભોમાં સહજ રીતે ભાષા ઢાળવાની ક્ષમતા વિકસાવવાનો છે."
    )

    sentences = [
        (
            f"Im heutigen Modul dient {k1} dazu, gedankliche Präzision sichtbar zu machen.",
            f"In today's module, {k1} helps make conceptual precision visible.",
            f"આજના મોડ્યુલમાં {k1} વિચારની ચોકસાઈને સ્પષ્ટ બનાવે છે."
        ),
        (
            f"Durch {k2} wirkt meine Argumentation differenzierter und belastbarer.",
            f"Through {k2}, my argumentation becomes more differentiated and robust.",
            f"{k2} દ્વારા મારી દલીલ વધુ સૂક્ષ્મ અને મજબૂત બને છે."
        ),
        (
            f"{k3} verschafft meinem Beitrag auf C2-Niveau mehr Tiefe und Reichweite.",
            f"{k3} gives my contribution more depth and reach at C2 level.",
            f"{k3} મારા પ્રતિભાવને C2 સ્તરે વધુ ઊંડાણ અને વ્યાપ આપે છે."
        ),
        (
            "Auf C2 zählt nicht nur, was gesagt wird, sondern auch mit welcher intellektuellen Feinsteuerung es geschieht.",
            "At C2, what matters is not only what is said, but also the intellectual fine control with which it is said.",
            "C2 પર ફક્ત શું કહેવામાં આવે છે તે જ નહીં, પણ કઈ બુદ્ધિપૂર્ણ સૂક્ષ્મ નિયંત્રણથી કહેવામાં આવે છે તે પણ મહત્વનું છે."
        ),
        (
            "Ich prüfe Register, Nuancen und implizite Wirkungen, bevor ich meine Endfassung festlege.",
            "I check register, nuances, and implicit effects before finalizing my final version.",
            "અંતિમ આવૃત્તિ નક્કી કરતાં પહેલાં હું રજિસ્ટર, સૂક્ષ્મતાઓ અને ગૂઢ અસર ચકાસું છું."
        ),
    ]

    tips = [
        ("📌", "C2 Focus", f"Go beyond accuracy: use {k1} and {k2} to create intellectual texture, not just correctness.", "E8F0FE"),
        ("💡", "Strategy", "Draft with freedom, then compress, refine, and rebalance tone until every sentence sounds intentional.", "FFF8E1"),
        ("🎯", "Progress", "C2 progress is visible when you can shift register without losing depth, elegance, or precision.", "E6F4EA"),
    ]

    tasks = [
        f"Write a 260–300 word expert-level text on '{subtopic}' using at least 8 target words with clear tonal control.",
        "Record a 4-minute spoken response that includes a counter-position, a refinement, and a strong closing synthesis.",
        "Rewrite your answer for a second audience (academic, diplomatic, or cultural) and note exactly how the register changes.",
    ]

    mini_test = [
        ("Translate into German: 'Precision at this level depends on nuance, structure, and controlled tone.'", "Präzision auf diesem Niveau hängt von Nuancen, Struktur und kontrolliertem Ton ab."),
        ("What is more C2-like: sounding complex at any cost or sounding exact, flexible, and context-sensitive?", "Sounding exact, flexible, and context-sensitive"),
        ("Provide one key word from today's vocabulary.", vocab[0][0]),
        ("At C2, what should improve together: vocabulary only or vocabulary + judgment + register control?", "Vocabulary + judgment + register control"),
        ("Name one final revision step before submission.", "Tighten nuance, tone, and implicit coherence"),
    ]

    outcome = (
        f"After Day {day_num}, you can approach '{subtopic}' with stronger C2 command, greater stylistic maturity, "
        f"and sharper intellectual control. Your speaking sounds more deliberate and your writing carries more authority."
    )

    return (day_num, title, topic, eng_exp, guj_exp, vocab, sentences, tips, tasks, mini_test, outcome)


def build_days():
    days = []
    for module_index, module in enumerate(MODULES):
        start, end = module["range"]
        for day in range(start, end + 1):
            i = day - 331
            subtopic = SUBTOPICS[i]
            idx = (day - start) + module_index * 3
            days.append(build_day_entry(day, module, subtopic, idx))
    return days


DAYS = build_days()


def build_book():
    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    add_para(doc, "🇩🇪", size_pt=64, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=32, space_after=8)
    add_para(doc, "DEUTSCH LERNEN", size_pt=34, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "C2 COURSE", size_pt=22, bold=True, color=PURPLE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "Days 331 to 365", size_pt=18, bold=True, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_para(doc, "ગુજરાતી ભાષકો માટે C2 જર્મન અભ્યાસ", size_pt=14, bold=True, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "For Gujarati Speakers | English + Gujarati Explanations", size_pt=11, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "35 Days · Precision · Interpretation · Diplomacy · Expert C2 Mastery", size_pt=11, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    add_para(doc, "📖 How to Use This C2 Book", size_pt=14, bold=True, color=BLUE, space_after=4)
    tips = [
        "📅 Study one day deeply; C2 gains come from slow, high-quality refinement.",
        "✍️ Write daily and cut weak phrasing without mercy.",
        "🗣️ Speak aloud for at least 20 minutes and listen for tone, register, and precision.",
        "🔁 Rework one old answer each week in a more elegant register.",
        "🎯 Focus on thought quality, not just vocabulary quantity.",
        "🧪 Treat the final days as integrated C2 simulations with strict timing.",
    ]
    for t in tips:
        add_bullet(doc, t)

    for entry in DAYS:
        add_day(doc, *entry)

    out = "German_Learning_Book_C2_Gujarati_Days_331_365.docx"
    doc.save(out)
    print(f"✅ Saved: {out}")
    print(f"   Days: {len(DAYS)}")


if __name__ == "__main__":
    build_book()
